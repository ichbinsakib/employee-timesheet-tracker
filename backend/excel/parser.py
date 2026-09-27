"""Header-based Excel timesheet parser.

Nothing here depends on fixed cell coordinates. The parser:

1. Scans each sheet for label cells ("Name", "Employee", "Date", ...) and takes
   the next non-empty cell to the right (or below) as the value.
2. Finds the table header row by scoring every row against known column
   synonyms (costing code, notes, file type, completed, features, hours).
3. Reads rows under the header until a "Total ..." row or the end of the sheet,
   skipping blank rows.
4. Captures a declared total (e.g. "Total Burden Time: 10.4") for validation.

It returns plain data; validation and storage happen elsewhere.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# canonical column -> list of header patterns (regex, matched against normalised header text)
COLUMN_PATTERNS: dict[str, list[str]] = {
    "costing_code": [r"costing\s*code", r"cost\s*code", r"project\s*(code|no|number|#)", r"job\s*(code|no|number|#)", r"charge\s*code", r"^code$"],
    "notes": [r"^notes?\b", r"description", r"task\s*details?", r"work\s*(done|performed|description)", r"activit(y|ies)", r"comments?"],
    "file_type": [r"file\s*type", r"^type$", r"document\s*type"],
    "completed": [r"completed", r"complete\?", r"^done\??$", r"status"],
    "feature_count": [r"features?", r"feature\s*count"],
    "burden_hours": [r"burden\s*hours?", r"burden\s*time", r"^hours?$", r"hours?\s*(spent|worked)", r"^hrs\.?$", r"time\s*spent", r"duration"],
}
REQUIRED_FOR_HEADER = 3  # a row needs at least this many recognised columns

NAME_LABELS = [r"^(employee\s*)?name\s*:?$", r"^employee\s*:?$", r"^prepared\s*by\s*:?$"]
DATE_LABELS = [r"^(work\s*)?date\s*:?$", r"^day\s*:?$", r"^date\s*of\s*work\s*:?$"]
TOTAL_PATTERN = re.compile(r"^total\b", re.I)
NA_VALUES = {"", "/", "-", "--", "n/a", "na", "none", "null", "x"}

DATE_FORMATS = ["%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%b %d, %Y", "%B %d, %Y", "%d.%m.%Y", "%Y/%m/%d"]


@dataclass
class ParsedRow:
    excel_row: int
    costing_code: str | None
    notes: str | None
    file_type: str | None
    completed_raw: str | None
    feature_raw: Any
    hours_raw: Any


@dataclass
class ParsedSheet:
    sheet_name: str
    employee_name: str | None = None
    work_date: date | None = None
    work_date_raw: Any = None
    header_row: int | None = None
    columns: dict[str, int] = field(default_factory=dict)  # canonical -> column index (1-based)
    rows: list[ParsedRow] = field(default_factory=list)
    declared_total: float | None = None
    declared_total_row: int | None = None


@dataclass
class ParseResult:
    sheets: list[ParsedSheet] = field(default_factory=list)
    skipped_sheets: list[str] = field(default_factory=list)
    error: str | None = None


# --------------------------------------------------------------------------- helpers
def norm_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", text).strip().lower()


def clean_str(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return None if text.lower() in NA_VALUES else text


def parse_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 < float(value) < 80000:  # Excel serial number
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    text = str(value).strip()
    text = re.sub(r"^(date\s*:?\s*)", "", text, flags=re.I)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().lower()
    if text in NA_VALUES:
        return None
    text = text.replace(",", ".") if text.count(",") == 1 and "." not in text else text.replace(",", "")
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(h|hr|hrs|hour|hours)?", text)
    if m:
        return float(m.group(1))
    m = re.fullmatch(r"(\d+):(\d{2})", text)  # 7:30
    if m:
        return int(m.group(1)) + int(m.group(2)) / 60
    return None


def date_from_filename(filename: str) -> date | None:
    stem = Path(filename).stem
    patterns = [
        (r"(\d{4})[-_.](\d{1,2})[-_.](\d{1,2})", ("y", "m", "d")),
        (r"(\d{1,2})[-_.](\d{1,2})[-_.](\d{4})", ("m", "d", "y")),
        (r"(\d{1,2})[-_.](\d{1,2})[-_.](\d{2})(?!\d)", ("m", "d", "y2")),
    ]
    for pattern, order in patterns:
        m = re.search(pattern, stem)
        if not m:
            continue
        parts = dict(zip(order, m.groups()))
        try:
            year = int(parts.get("y") or 2000 + int(parts["y2"]))
            return date(year, int(parts["m"]), int(parts["d"]))
        except ValueError:
            continue
    return None


# --------------------------------------------------------------------------- core
def _match_column(header: str) -> str | None:
    if not header:
        return None
    for canonical, patterns in COLUMN_PATTERNS.items():
        for p in patterns:
            if re.search(p, header):
                return canonical
    return None


def _find_header(grid: list[list[Any]]) -> tuple[int | None, dict[str, int]]:
    best_row, best_cols = None, {}
    for r_idx, row in enumerate(grid[:60]):
        cols: dict[str, int] = {}
        for c_idx, value in enumerate(row):
            canonical = _match_column(norm_text(value))
            if canonical and canonical not in cols:
                cols[canonical] = c_idx
        if len(cols) > len(best_cols):
            best_row, best_cols = r_idx, cols
    if len(best_cols) >= REQUIRED_FOR_HEADER and "burden_hours" in best_cols:
        return best_row, best_cols
    return None, {}


TITLE_WORDS = re.compile(r"\b(sheet|timesheet|form|report|production|template)\b", re.I)
ALL_LABELS = NAME_LABELS + DATE_LABELS


def _is_label_or_title(value: Any) -> bool:
    text = norm_text(value)
    return bool(TITLE_WORDS.search(text)) or any(re.search(p, text) for p in ALL_LABELS)


def _labelled_value(grid: list[list[Any]], label_patterns: list[str], stop_row: int | None, kind: str) -> Any:
    """Value next to (right of, else below) the first cell matching a label.

    Only looks a few cells to the right and never returns another label or a
    sheet title, so an empty "Name" cell does not pick up "Employee Production Sheet".
    """
    def acceptable(v: Any) -> bool:
        if v is None or not str(v).strip() or _is_label_or_title(v) or _match_column(norm_text(v)):
            return False
        if kind == "date":
            return parse_date(v) is not None
        return isinstance(v, str) and parse_date(v) is None and bool(re.search(r"[a-zA-Z]", v))

    limit = stop_row if stop_row is not None else min(len(grid), 20)
    for r_idx in range(limit):
        row = grid[r_idx]
        for c_idx, value in enumerate(row):
            text = norm_text(value)
            if not text or not any(re.search(p, text) for p in label_patterns):
                continue
            neighbours = [v for v in row[c_idx + 1:c_idx + 4] if v is not None and str(v).strip()]
            if neighbours and acceptable(neighbours[0]):
                return neighbours[0]
            if r_idx + 1 < limit and c_idx < len(grid[r_idx + 1]):
                below = grid[r_idx + 1][c_idx]
                if acceptable(below):
                    return below
            if kind == "date" and neighbours and not _is_label_or_title(neighbours[0]):
                return neighbours[0]  # unreadable date: return raw so validation can report it
    # "Name: John Smith" in a single cell
    for r_idx in range(limit):
        for value in grid[r_idx]:
            if isinstance(value, str) and ":" in value:
                label, _, rest = value.partition(":")
                if rest.strip() and any(re.search(p, norm_text(label)) for p in label_patterns):
                    return rest.strip()
    return None


def parse_sheet(ws) -> ParsedSheet | None:
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    header_idx, cols = _find_header(grid)
    if header_idx is None:
        return None

    sheet = ParsedSheet(sheet_name=ws.title, header_row=header_idx + 1, columns={k: v + 1 for k, v in cols.items()})
    name_val = _labelled_value(grid, NAME_LABELS, header_idx, "name")
    sheet.employee_name = clean_str(name_val)
    date_val = _labelled_value(grid, DATE_LABELS, header_idx, "date")
    sheet.work_date_raw = date_val
    sheet.work_date = parse_date(date_val)

    def cell(row: list[Any], key: str) -> Any:
        idx = cols.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    for r_idx in range(header_idx + 1, len(grid)):
        row = grid[r_idx]
        texts = [norm_text(v) for v in row]
        if not any(texts):
            continue
        total_label = next((i for i, t in enumerate(texts) if TOTAL_PATTERN.match(t)), None)
        if total_label is not None:
            # Declared total: the hours column on this row, else the first number after the label
            total = parse_number(cell(row, "burden_hours"))
            if total is None:
                total = next((parse_number(v) for v in row[total_label + 1:] if parse_number(v) is not None), None)
            sheet.declared_total, sheet.declared_total_row = total, r_idx + 1
            break
        parsed = ParsedRow(
            excel_row=r_idx + 1,
            costing_code=clean_str(cell(row, "costing_code")),
            notes=clean_str(cell(row, "notes")),
            file_type=clean_str(cell(row, "file_type")),
            completed_raw=clean_str(cell(row, "completed")),
            feature_raw=cell(row, "feature_count"),
            hours_raw=cell(row, "burden_hours"),
        )
        if parsed.costing_code or parsed.notes or parse_number(parsed.hours_raw) is not None:
            sheet.rows.append(parsed)
    return sheet


def parse_workbook(content: bytes | Path | str) -> ParseResult:
    result = ParseResult()
    try:
        source = BytesIO(content) if isinstance(content, (bytes, bytearray)) else content
        # data_only=True returns the values Excel last calculated for formulas
        wb = load_workbook(source, data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001 - any openpyxl failure means "not a readable xlsx"
        result.error = f"Could not open the file as an Excel workbook ({type(exc).__name__}: {exc})"
        return result
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        sheet = parse_sheet(ws)
        if sheet is None:
            result.skipped_sheets.append(ws.title)
        else:
            result.sheets.append(sheet)
    if not result.sheets:
        result.error = "No timesheet table found. Expected a header row with columns such as Costing Code, Notes and Burden Hours."
    return result
