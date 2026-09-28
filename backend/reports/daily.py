"""Daily management report: build the data once, export as Excel, CSV or PDF."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime, timedelta
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.analytics import alerts as alerts_mod
from backend.analytics import kpi
from backend.config import settings
from backend.models import Employee, GeneratedReport, TimesheetEntry, now
from backend.services import costing_codes as cc

FORMATS = ("xlsx", "csv", "pdf")


def build_daily_report(db: Session, day: date, now_dt: datetime | None = None) -> dict:
    now_dt = now_dt or datetime.now()
    cats = kpi.codes(db)
    status = kpi.day_status(db, day, now=now_dt)
    ds = kpi.load(db, day, day)
    today = kpi.summarize(ds, cats)

    hist_start, hist_end = day - timedelta(days=30), day - timedelta(days=1)
    hist = kpi.load(db, hist_start, hist_end)
    hist_summary = kpi.summarize(hist, cats)

    employees = []
    submitted_ids = {e.employee_id for e in ds.entries}
    for emp in db.scalars(select(Employee).where(Employee.active.is_(True)).order_by(Employee.name)):
        s = kpi.summarize(ds, cats, emp.id)
        h = kpi.summarize(hist, cats, emp.id)
        employees.append({
            "employee_id": emp.id, "employee": emp.name,
            "submitted": emp.id in submitted_ids,
            "hours": s["total_hours"], "tasks": s["task_count"], "completed": s["completed_tasks"],
            "incomplete": s["incomplete_tasks"], "completion_rate": s["completion_rate"],
            "main_codes": ", ".join(f"{c['code']} {c['percent']:g}%" for c in s["codes"][:3] if c["percent"]),
            "avg_daily_hours_30d": h["avg_daily_hours"],
            "hours_vs_30d": round(s["total_hours"] - h["avg_daily_hours"], 2) if emp.id in submitted_ids and h["avg_daily_hours"] else None,
        })

    changes = []
    if hist_summary["total_hours"]:
        prev = {c["code"]: c["percent"] or 0 for c in hist_summary["codes"]}
        cur = {c["code"]: c["percent"] or 0 for c in today["codes"]}
        labels = {c["code"]: c["label"] for c in hist_summary["codes"] + today["codes"]}
        for name in sorted(set(prev) | set(cur)):
            changes.append({"code": name, "label": labels.get(name, name), "today_pct": cur.get(name, 0), "last_30d_pct": prev.get(name, 0),
                            "change": round(cur.get(name, 0) - prev.get(name, 0), 1)})
        changes.sort(key=lambda r: -abs(r["change"]))

    attention = alerts_mod.build_alerts(db, day, lookback_days=1, now=now_dt)
    entries = db.execute(
        select(TimesheetEntry, Employee.name)
        .join(Employee, Employee.id == TimesheetEntry.employee_id)
        .where(TimesheetEntry.work_date == day).order_by(Employee.name, TimesheetEntry.original_excel_row)
    ).all()

    return {
        "date": day.isoformat(),
        "generated_at": now_dt.isoformat(timespec="seconds"),
        "status": status,
        "summary": {
            "employees_expected": status["expected"], "employees_submitted": status["submitted"],
            "missing": len(status["missing"]), "pending": len(status["pending"]),
            "total_hours": today["total_hours"], "tasks": today["task_count"], "activities": today["activity_count"],
            "completed": today["completed_tasks"], "incomplete": today["incomplete_tasks"],
            "completion_not_recorded": today["completion_not_recorded"], "completion_rate": today["completion_rate"],
            "avg_daily_hours_30d": hist_summary["avg_daily_hours"],
        },
        "codes": today["codes"],
        "code_changes": changes,
        "employees": employees,
        "anomalies": kpi.daily_hours_anomalies(db, day, day),
        "data_quality": kpi.data_quality(db, day, day, min_severity="warning"),
        "attention": attention,
        "entries": [
            {"employee": name, "costing_code": e.costing_code,
             "code_description": cats.get((e.costing_code or "").strip().upper()), "title": e.title, "notes": e.notes,
             "file_type": e.file_type, "completed": {True: "Yes", False: "No"}.get(e.completed, "Not recorded"),
             "features": e.feature_count, "hours": e.burden_hours}
            for e, name in entries
        ],
    }


def _fmt(v) -> str:
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:g}"
    return str(v)


# --------------------------------------------------------------------------- Excel
def to_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3A5F")

    def sheet(title: str, headers: list[str], rows: list[list], widths: list[int] | None = None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [18] * len(headers), start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    ws = wb.active
    ws.title = "Summary"
    ws.append([f"Daily Timesheet Report — {report['date']}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {report['generated_at']}"])
    ws.append([])
    labels = {
        "employees_expected": "Employees expected", "employees_submitted": "Timesheets received",
        "missing": "Missing (deadline passed)", "pending": "Not yet due", "total_hours": "Total hours",
        "tasks": "Tasks (rows)", "activities": "Activities (note lines)", "completed": "Completed tasks",
        "incomplete": "Incomplete tasks", "completion_not_recorded": "Completion not recorded",
        "completion_rate": "Completion rate (%)", "avg_daily_hours_30d": "Avg daily hours per employee, last 30 days",
    }
    for k, label in labels.items():
        ws.append([label, report["summary"].get(k)])
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16

    sheet("Employees", ["Employee", "Submitted", "Hours", "Tasks", "Completed", "Incomplete", "Completion %", "Main costing codes", "Avg daily hours (30d)", "Hours vs 30d avg"],
          [[e["employee"], "Yes" if e["submitted"] else "No", e["hours"], e["tasks"], e["completed"], e["incomplete"], e["completion_rate"],
            e["main_codes"], e["avg_daily_hours_30d"], e["hours_vs_30d"]] for e in report["employees"]],
          [26, 10, 8, 8, 10, 10, 12, 50, 18, 16])
    sheet("Costing codes", ["Code", "Description", "Hours", "Tasks", "Percent"],
          [[c["code"], c["description"], c["hours"], c["tasks"], c["percent"]] for c in report["codes"]], [12, 50, 10, 8, 10])
    sheet("Changes vs 30 days", ["Costing code", "Today %", "Last 30 days %", "Change (points)"],
          [[c["label"], c["today_pct"], c["last_30d_pct"], c["change"]] for c in report["code_changes"]], [50, 12, 16, 16])
    sheet("Attention", ["Type", "Detail"], [[a["title"], a["message"]] for a in report["attention"]], [28, 110])
    sheet("Data quality", ["Employee", "Severity", "Issue", "Excel row", "File"],
          [[q["employee"], q["severity"], q["message"], q["excel_row"], q["file"]] for q in report["data_quality"]], [24, 10, 80, 10, 40])
    ws_e = sheet("Entries", ["Employee", "Costing code", "Code description", "Notes", "File type", "Completed", "Features", "Hours"],
                 [[e["employee"], e["costing_code"], e["code_description"], e["notes"], e["file_type"], e["completed"], e["features"], e["hours"]] for e in report["entries"]],
                 [24, 14, 22, 80, 14, 12, 10, 8])
    for row in ws_e.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- CSV
def to_csv(report: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([f"Daily Timesheet Report {report['date']}", f"Generated {report['generated_at']}"])
    w.writerow([])
    w.writerow(["Summary"])
    for k, v in report["summary"].items():
        w.writerow([k, v])
    w.writerow([])
    w.writerow(["Employee", "Submitted", "Hours", "Tasks", "Completed", "Incomplete", "Completion %", "Main costing codes", "Avg daily hours (30d)", "Hours vs 30d avg"])
    for e in report["employees"]:
        w.writerow([e["employee"], "Yes" if e["submitted"] else "No", e["hours"], e["tasks"], e["completed"], e["incomplete"],
                    e["completion_rate"], e["main_codes"], e["avg_daily_hours_30d"], e["hours_vs_30d"]])
    w.writerow([])
    w.writerow(["Costing code", "Description", "Hours", "Tasks", "Percent"])
    for c in report["codes"]:
        w.writerow([c["code"], c["description"], c["hours"], c["tasks"], c["percent"]])
    w.writerow([])
    w.writerow(["Attention", "Detail"])
    for a in report["attention"]:
        w.writerow([a["title"], a["message"]])
    w.writerow([])
    w.writerow(["Data quality: employee", "Severity", "Issue", "Excel row", "File"])
    for q in report["data_quality"]:
        w.writerow([q["employee"], q["severity"], q["message"], q["excel_row"], q["file"]])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


# --------------------------------------------------------------------------- PDF
def to_pdf(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=8, leading=10)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"Daily Timesheet Report {report['date']}")
    story = [Paragraph(f"Daily Timesheet Report — {report['date']}", styles["Title"]),
             Paragraph(f"Generated {report['generated_at']}. Figures describe what was recorded in timesheets, grouped by costing code.", small),
             Spacer(1, 6)]

    def table(headers, rows, widths=None):
        data = [[Paragraph(f"<b>{h}</b>", small) for h in headers]] + [[Paragraph(_esc(_fmt(c)), small) for c in r] for r in rows]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5EAF1")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C2CF")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return t

    s = report["summary"]
    story += [Paragraph("Summary", styles["Heading2"]), table(
        ["Expected", "Received", "Missing", "Not yet due", "Total hours", "Tasks", "Completed", "Incomplete", "Not recorded", "Completion %"],
        [[s["employees_expected"], s["employees_submitted"], s["missing"], s["pending"], s["total_hours"], s["tasks"], s["completed"],
          s["incomplete"], s["completion_not_recorded"], s["completion_rate"]]])]
    if report["attention"]:
        story += [Paragraph("Management attention", styles["Heading2"]),
                  table(["Type", "Detail"], [[a["title"], a["message"]] for a in report["attention"]], [45 * mm, 225 * mm])]
    story += [Paragraph("Employees", styles["Heading2"]), table(
        ["Employee", "Submitted", "Hours", "Tasks", "Done", "Not done", "Main costing codes", "Avg/day 30d", "vs 30d"],
        [[e["employee"], "Yes" if e["submitted"] else "No", e["hours"], e["tasks"], e["completed"], e["incomplete"],
          e["main_codes"], e["avg_daily_hours_30d"], e["hours_vs_30d"]] for e in report["employees"]],
        [45 * mm, 18 * mm, 15 * mm, 13 * mm, 13 * mm, 16 * mm, 95 * mm, 22 * mm, 18 * mm])]
    if report["codes"]:
        story += [Paragraph("Hours by costing code", styles["Heading2"]),
                  table(["Code", "Description", "Hours", "Tasks", "%"],
                        [[c["code"], c["description"], c["hours"], c["tasks"], c["percent"]] for c in report["codes"]],
                        [22 * mm, 150 * mm, 20 * mm, 15 * mm, 15 * mm])]
    if report["code_changes"]:
        story += [Paragraph("Change vs last 30 days", styles["Heading2"]),
                  table(["Costing code", "Today %", "Last 30 days %", "Change (pts)"],
                        [[c["label"], c["today_pct"], c["last_30d_pct"], c["change"]] for c in report["code_changes"][:8]],
                        [120 * mm, 25 * mm, 30 * mm, 25 * mm])]
    if report["anomalies"]:
        story += [Paragraph("Unusual hours", styles["Heading2"]),
                  table(["Employee", "Hours", "Recent average", "Note"],
                        [[a["employee"], a["hours"], a["recent_average"], a["reason"]] for a in report["anomalies"]])]
    if report["data_quality"]:
        story += [Paragraph("Data quality", styles["Heading2"]),
                  table(["Employee", "Severity", "Issue", "Row"], [[q["employee"], q["severity"], q["message"], q["excel_row"]] for q in report["data_quality"]],
                        [45 * mm, 20 * mm, 185 * mm, 15 * mm])]
    doc.build(story)
    return buf.getvalue()


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")


EXPORTERS = {"xlsx": to_xlsx, "csv": to_csv, "pdf": to_pdf}
MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "pdf": "application/pdf",
}


def generate_and_save(db: Session, day: date, formats=FORMATS) -> list[Path]:
    report = build_daily_report(db, day)
    folder = settings.reports_dir / f"{day:%Y}" / f"{day:%m}"
    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    for fmt in formats:
        path = folder / f"daily_report_{day.isoformat()}.{fmt}"
        path.write_bytes(EXPORTERS[fmt](report))
        row = db.scalar(select(GeneratedReport).where(GeneratedReport.report_date == day, GeneratedReport.fmt == fmt))
        if row:
            row.path, row.created_at = str(path), now()
        else:
            db.add(GeneratedReport(report_date=day, fmt=fmt, path=str(path)))
        paths.append(path)
    db.commit()
    return paths
