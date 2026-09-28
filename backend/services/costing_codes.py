"""Costing-code list: read from the "COSTING CODE" sheet that CAMCO timesheets carry.

Work is reported by costing code (the company's own structure) rather than by
guessed keyword categories.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.models import CostingCode, Submission, TimesheetEntry, now

CODE_HEADERS = {"code", "costing code", "cost code", "costing"}
DESC_HEADERS = {"details", "detail", "description", "desc", "name", "project"}
CODE_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._\-/]{0,29}$")
NO_CODE = "(none)"


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def clean_description(v) -> str | None:
    s = re.sub(r"\s+", " ", str(v or "")).strip().rstrip("-–— ").strip()
    return s or None


def extract_codes(wb) -> dict[str, str | None]:
    """Find sheets (hidden or not) whose header row has a code column and a details column."""
    out: dict[str, str | None] = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True, max_row=5000))
        for h_idx, row in enumerate(rows[:5]):
            heads = [_norm(c) for c in row]
            code_col = next((i for i, h in enumerate(heads) if h in CODE_HEADERS), None)
            desc_col = next((i for i, h in enumerate(heads) if h in DESC_HEADERS), None)
            if code_col is None or desc_col is None:
                continue
            for r in rows[h_idx + 1:]:
                code = str(r[code_col]).strip() if code_col < len(r) and r[code_col] is not None else ""
                if not code or not CODE_RE.match(code):
                    continue
                desc = clean_description(r[desc_col]) if desc_col < len(r) else None
                out[code.upper()] = desc
            break
    return out


def extract_codes_from_bytes(content: bytes) -> dict[str, str | None]:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - not a readable workbook: nothing to learn
        return {}
    try:
        return extract_codes(wb)
    finally:
        wb.close()


def upsert(db: Session, codes: dict[str, str | None]) -> int:
    """Add new codes and refresh descriptions. Descriptions edited by hand are kept."""
    changed = 0
    for code, desc in codes.items():
        row = db.get(CostingCode, code)
        if row is None:
            db.add(CostingCode(code=code, description=desc, source="timesheet"))
            changed += 1
        elif row.source != "manual" and desc and row.description != desc:
            row.description, row.updated_at = desc, now()
            changed += 1
    return changed


def backfill_from_saved_files(db: Session) -> int:
    """Learn codes from timesheet files already received (used once after upgrading)."""
    total = 0
    paths = {s.stored_path for s in db.scalars(select(Submission).where(Submission.stored_path.is_not(None)))}
    for p in sorted(paths):
        if p and Path(p).exists():
            total += upsert(db, extract_codes_from_bytes(Path(p).read_bytes()))
    return total


class CodeBook(dict):
    """code -> description. `listed` holds the codes that are in the COSTING CODE list;
    other codes get a description taken from their most recent timesheet note."""

    def __init__(self, *args, listed: set[str] | None = None, **kw):
        super().__init__(*args, **kw)
        self.listed = listed or set()


def note_summary(notes: str | None) -> str | None:
    """First line of a note, without trailing dashes, as a short description."""
    if not notes:
        return None
    first = next((ln.strip() for ln in notes.splitlines() if ln.strip()), "")
    first = re.sub(r"^[-–—•*]\s*", "", first).rstrip(" -–—").strip()
    if not first:
        return None
    return first if len(first) <= 80 else first[:77].rstrip() + "…"


def descriptions(db: Session) -> CodeBook:
    book = CodeBook({c.code: c.description for c in db.scalars(select(CostingCode))})
    book.listed = set(book)
    # codes used on timesheets but missing from the list: describe them from the latest note
    rows = db.execute(
        select(TimesheetEntry.costing_code, TimesheetEntry.notes)
        .where(TimesheetEntry.costing_code.is_not(None), TimesheetEntry.notes.is_not(None))
        .order_by(TimesheetEntry.work_date.desc(), TimesheetEntry.id.desc())
    )
    for code, notes in rows:
        key = (code or "").strip().upper()
        if key and key not in book:
            summary = note_summary(notes)
            if summary:
                book[key] = f"{summary} (from notes)"
    return book


def label(code: str | None, desc: dict[str, str | None]) -> str:
    if not code:
        return "No costing code"
    d = desc.get(code.upper())
    return f"{code} – {d}" if d else code


def all_codes_with_usage(db: Session) -> list[dict]:
    from sqlalchemy import func
    usage = {r[0].upper(): (r[1], r[2]) for r in db.execute(
        select(TimesheetEntry.costing_code, func.count(), func.sum(TimesheetEntry.burden_hours))
        .where(TimesheetEntry.costing_code.is_not(None)).group_by(TimesheetEntry.costing_code))}
    rows = []
    for c in db.scalars(select(CostingCode).order_by(CostingCode.code)):
        n, h = usage.pop(c.code, (0, 0))
        rows.append({"code": c.code, "description": c.description, "source": c.source, "entries": n, "hours": round(h or 0, 2), "in_list": True})
    book = descriptions(db)
    for code, (n, h) in sorted(usage.items()):  # used on timesheets but not in any list
        rows.append({"code": code, "description": None, "from_notes": book.get(code), "source": None,
                     "entries": n, "hours": round(h or 0, 2), "in_list": False})
    return rows
