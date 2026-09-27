"""Create Excel timesheets in the same layout as the real "Employee Production Sheet".

Used by the tests and for trying the importer by hand:

    python scripts/make_sample_timesheet.py --out samples/
"""
from __future__ import annotations

import argparse
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REAL_EXAMPLE_NOTES = "\n".join([
    "Co-ordination task",
    "-SO Release",
    "-Performing SO release-related task.",
    "-Future CO shortage data analysis.",
    "-Taking followup for futuer CO Shortage",
    "-Update SO tracker sheet",
    "-Master Shudule sheet update for CO shortage",
    "- Kanban checking",
    "- Making Subcontract PO",
    "-Teams & Email comunication for followup",
    "-Make IT ticket for Days Behind App",
    "-Communication with production",
    "-Checking the RMA sheet for Raymond.",
    "-Approve the invoice.",
    "-Through Purchasing 2, get vendor quotations and check them.",
    "-Entering data in the Matrix data entry sheet.",
    "-Coordination-related other tasks.",
])

HEADERS = [
    "COSTING CODE",
    "Notes (BE AS DETAILED AND SPECIFIC AS POSSIBLE!!!)",
    "File Type (Drawing, Assembly Drawing, Model, Assembly)",
    "Completed? (Yes/No)",
    "Total # of Features on a File",
    "Burden Hours",
]


def build_timesheet(
    path: Path | str,
    name: str = "Md. Kamrul Hasan",
    work_date: date = date(2026, 9, 23),
    rows: list[tuple] | None = None,
    blank_rows: int = 4,
    total: float | str | None = "sum",
    offset_row: int = 1,
    offset_col: int = 1,
) -> Path:
    """rows: (costing_code, notes, file_type, completed, features, hours)."""
    if rows is None:
        rows = [("P4627", REAL_EXAMPLE_NOTES, "/", "/", "/", 10.4)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    r0, c0 = offset_row, offset_col
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    grey = PatternFill("solid", fgColor="D9D9D9")

    ws.cell(r0, c0, "Name")
    ws.cell(r0, c0 + 1, name)
    ws.cell(r0, c0 + 2, "Employee Production Sheet")
    ws.merge_cells(start_row=r0, start_column=c0 + 2, end_row=r0, end_column=c0 + 3)
    ws.cell(r0, c0 + 4, "Date")
    ws.cell(r0, c0 + 5, datetime(work_date.year, work_date.month, work_date.day)).number_format = "m/d/yyyy"

    for i, h in enumerate(HEADERS):
        cell = ws.cell(r0 + 1, c0 + i, h)
        cell.font = Font(bold=True)
        cell.fill = grey
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = border

    r = r0 + 2
    for row in rows:
        for i, value in enumerate(row):
            cell = ws.cell(r, c0 + i, value)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = border
        r += 1
    r += blank_rows
    ws.cell(r, c0 + 1, "Total Burden Time:").font = Font(bold=True)
    if total == "sum":
        total = round(sum(float(x[5]) for x in rows if isinstance(x[5], (int, float))), 2)
    if total is not None:
        ws.cell(r, c0 + 5, total).font = Font(bold=True)

    for col, width in zip("ABCDEF", (14, 70, 22, 14, 14, 12)):
        ws.column_dimensions[col].width = width
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="samples")
    args = ap.parse_args()
    out = Path(args.out)
    p = build_timesheet(out / "Kamrul_Hasan_2026-09-23.xlsx")
    print(f"Wrote {p}")
