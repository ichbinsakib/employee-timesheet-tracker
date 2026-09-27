from datetime import date

from backend.excel.parser import parse_workbook, parse_number, date_from_filename
from backend.services.classifier import split_notes


def test_parses_real_example_layout(tmp_xlsx):
    path = tmp_xlsx()
    res = parse_workbook(path.read_bytes())
    assert res.error is None
    sheet = res.sheets[0]
    assert sheet.employee_name == "Md. Kamrul Hasan"
    assert sheet.work_date == date(2026, 9, 23)
    assert set(sheet.columns) == {"costing_code", "notes", "file_type", "completed", "feature_count", "burden_hours"}
    assert len(sheet.rows) == 1
    row = sheet.rows[0]
    assert row.costing_code == "P4627"
    assert row.file_type is None and row.completed_raw is None  # "/" means not applicable
    assert row.hours_raw == 10.4
    assert sheet.declared_total == 10.4


def test_layout_shifted_and_columns_reordered(tmp_path):
    """Header-based detection: moving the table must not break parsing."""
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws["C3"] = "Employee Name:"; ws["D3"] = "Jane Roe"
    ws["C4"] = "Date"; ws["D4"] = "09/24/2026"
    headers = ["Hours", "Project Code", "Description", "Done?", "Features", "File Type"]
    for i, h in enumerate(headers):
        ws.cell(7, 2 + i, h)
    ws.cell(8, 2, "7.5"); ws.cell(8, 3, "P1"); ws.cell(8, 4, "Drawing updates for bracket"); ws.cell(8, 5, "Yes"); ws.cell(8, 6, 12); ws.cell(8, 7, "Drawing")
    ws.cell(10, 2, 0.5); ws.cell(10, 4, "Team meeting"); ws.cell(10, 5, "no")
    ws.cell(12, 3, "TOTAL"); ws.cell(12, 2, 8)
    p = tmp_path / "x.xlsx"; wb.save(p)
    sheet = parse_workbook(p.read_bytes()).sheets[0]
    assert sheet.employee_name == "Jane Roe"
    assert sheet.work_date == date(2026, 9, 24)
    assert [r.notes for r in sheet.rows] == ["Drawing updates for bracket", "Team meeting"]
    assert sheet.declared_total == 8


def test_not_a_timesheet(tmp_path):
    from openpyxl import Workbook
    wb = Workbook(); wb.active["A1"] = "hello"
    p = tmp_path / "x.xlsx"; wb.save(p)
    assert parse_workbook(p.read_bytes()).error
    assert parse_workbook(b"not excel").error


def test_helpers():
    assert parse_number("7:30") == 7.5
    assert parse_number("8 hrs") == 8
    assert parse_number("/") is None
    assert date_from_filename("Timesheet_2026-09-23.xlsx") == date(2026, 9, 23)
    assert date_from_filename("ts 9-23-2026.xlsx") == date(2026, 9, 23)


def test_split_notes():
    s = split_notes("Co-ordination task\n-SO Release\n- Kanban checking\n  continued")
    assert s.title == "Co-ordination task"
    assert s.activities == ["SO Release", "Kanban checking continued"]
    assert split_notes("Just one thing").activities == ["Just one thing"]
