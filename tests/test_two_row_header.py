"""Regression: the real CAMCO .xlsm splits headings over two rows ("COSTING"/"CODE",
"Burden"/"Hours") and carries hidden legacy sheets plus a visible lookup sheet."""
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import select

from backend.excel.parser import parse_workbook
from backend.models import Submission, TimesheetEntry
from backend.services.importer import ImportContext, import_file, retry_submission


def build_camco_style(path: Path) -> Path:
    wb = Workbook()
    legacy = wb.active
    legacy.title = "CX878-3 REV-B"
    legacy.sheet_state = "hidden"
    legacy.append([None, "Production Time", None, "DATE:", None, None, None, None, "EMPLOYEE:"])
    legacy.append([])
    legacy.append(["DEPT. CODE", "Shop Order #", "Part Number", None, None, "OP#", "Bad Pcs", "Good Pcs", "$ Per Piece",
                   "Total Dollars", "Routing", "Actual", "Total Minutes", "Setup Time (Hrs.)", "Burden Time (Hrs.)", "Prod. Time (Hrs.)"])
    ws = wb.create_sheet("DPS")
    ws["A1"], ws["B1"], ws["E1"], ws["I1"], ws["J1"] = "Name", "Test Person", "Production Sheet", "Date", datetime(2026, 9, 25)
    ws["A2"], ws["B2"], ws["G2"], ws["H2"], ws["I2"], ws["J2"] = (
        "COSTING", "Notes (BE AS DETAILED AND SPECIFIC AS POSSIBLE!!!)",
        "File Type (Drawing, Assembly Drawing, Model, Assembly", "Completed?", "Total # of Features on a File", "Burden")
    ws["A3"], ws["H3"], ws["J3"] = "CODE", "(Yes/No)", "Hours"
    rows = [("P3096", "Complete daily entry of employee times", 4.3), ("R6538", "Daily inventory count", 0.4),
            ("P1169", "Timesheet vs tracker review", 0), ("P3137", "Prepare monthly attendance sheet", 0.7)]
    for i, (code, note, hours) in enumerate(rows, start=5):
        ws.cell(i, 1, code); ws.cell(i, 2, note)
        for col in (7, 8, 9):
            ws.cell(i, col, "/")
        ws.cell(i, 10, hours)
    ws["B16"], ws["J16"] = "Total Burden Time:", 5.4000000000000004
    ws["A17"] = "How to count features:"
    ws["A18"] = "Drawings & Assembly Drawings: Count every dimension and note"
    lookup = wb.create_sheet("COSTING CODE")
    lookup.append(["CODE ", "DETAILS "])
    lookup.append(["P0205", "MANUFACTURING CHECKSHEET ENTRY"])
    wb.save(path)
    return path


def test_two_row_headers_are_joined(tmp_path):
    res = parse_workbook(build_camco_style(tmp_path / "09-25-2026 Test Person.xlsm").read_bytes())
    assert res.error is None and [s.sheet_name for s in res.sheets] == ["DPS"]
    sheet = res.sheets[0]
    assert set(sheet.columns) == {"costing_code", "notes", "file_type", "completed", "feature_count", "burden_hours"}
    assert sheet.employee_name == "Test Person" and str(sheet.work_date) == "2026-09-25"
    assert [r.costing_code for r in sheet.rows] == ["P3096", "R6538", "P1169", "P3137"]
    assert abs(sheet.declared_total - 5.4) < 1e-9  # the footer notes after the total are not rows


def test_retry_failed_import(db, tmp_path):
    path = build_camco_style(tmp_path / "t.xlsm")
    failed = Submission(source="gmail", attachment_filename="t.xlsm", file_sha256="x" * 64, stored_path=str(path),
                        import_status="failed", error_message="No timesheet table found.", gmail_message_id="m1")
    db.add(failed)
    db.commit()
    new = retry_submission(db, failed)
    db.commit()
    assert [s.import_status for s in new] == ["imported"]
    assert new[0].gmail_message_id == "m1" and new[0].total_hours == 5.4
    assert db.query(Submission).count() == 1 and db.query(TimesheetEntry).count() == 4


def test_costing_code_list_is_learned_on_import(db, tmp_path):
    from backend.analytics import kpi
    from backend.models import CostingCode
    from backend.services import costing_codes as cc
    path = build_camco_style(tmp_path / "09-25-2026 Test Person.xlsm")
    import_file(db, path.read_bytes(), ImportContext(filename=path.name))
    db.commit()
    assert db.get(CostingCode, "P0205").description == "MANUFACTURING CHECKSHEET ENTRY"
    # a hand-edited description is not overwritten by later timesheets
    row = db.get(CostingCode, "P0205")
    row.description, row.source = "Checksheet entry (edited)", "manual"
    db.commit()
    cc.upsert(db, {"P0205": "MANUFACTURING CHECKSHEET ENTRY", "P9999": "NEW CODE"})
    db.commit()
    assert db.get(CostingCode, "P0205").description == "Checksheet entry (edited)"
    assert db.get(CostingCode, "P9999").description == "NEW CODE"
    # codes used on rows but not in the list are still reported, flagged as not listed
    day = date(2026, 9, 25)
    s = kpi.summarize(kpi.load(db, day, day), kpi.codes(db))
    by_code = {c["code"]: c for c in s["codes"]}
    assert by_code["P3096"]["hours"] == 4.3 and by_code["P3096"]["in_list"] is False
    assert "P1169" not in by_code  # 0-hour row
